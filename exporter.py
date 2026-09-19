"""Render a project dict to JSON, DOCX or PDF."""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from schema import PRODUCT_EN, PRODUCT_ZH, steps_for


def _lang(node: dict, lang: str) -> str:
    if not isinstance(node, dict):
        return str(node or "")
    return node.get(lang) or node.get("zh") or node.get("en") or ""


def _field_map(step: dict) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for f in step.get("fields") or []:
        out[f["key"]] = f
    for sub in step.get("subs") or []:
        for f in sub.get("fields") or []:
            out[f["key"]] = f
    return out


def _option_label(field: dict, value: str, lang: str, extras: list | None = None) -> str:
    for opt in list(field.get("options") or []) + list(extras or []):
        if opt.get("id") == value:
            return _lang(opt, lang)
    return value


def _display(field: dict, value: Any, lang: str, custom: dict | None = None) -> str:
    if value is None or value == "":
        return "—"
    extras = (custom or {}).get(field.get("key") or "") or []
    kind = field.get("kind")
    if kind in ("radio", "select") and isinstance(value, str):
        return _option_label(field, value, lang, extras)
    if kind == "check" and isinstance(value, list):
        return "、".join(_option_label(field, v, lang, extras) for v in value) or "—"
    return str(value)


def iter_sections(data: dict, scope: str, step_id: str | None, sub_id: str | None, lang: str):
    doc_type = (data.get("docType") or "game")
    steps = steps_for(doc_type)
    custom = data.get("__customOptions") or {}
    if scope == "step":
        steps = [s for s in steps if s["id"] == step_id]
    elif scope == "sub":
        steps = [s for s in steps if s["id"] == step_id]
    for step in steps:
        title = f"{step['no']}  {_lang(step['title'], lang)}"
        lead = _lang(step.get("lead") or {}, lang)
        if scope == "sub" and step.get("subs"):
            for sub in step["subs"]:
                if sub["id"] != sub_id:
                    continue
                yield {
                    "title": f"{sub['no']}  {_lang(sub['title'], lang)}",
                    "lead": "",
                    "rows": [
                        {
                            "label": _lang(f["label"], lang),
                            "hint": _lang(f["hint"], lang),
                            "value": _display(f, data.get(f["key"]), lang, custom),
                        }
                        for f in sub.get("fields") or []
                    ],
                }
            return
        rows = []
        for f in step.get("fields") or []:
            rows.append(
                {
                    "label": _lang(f["label"], lang),
                    "hint": _lang(f["hint"], lang),
                    "value": _display(f, data.get(f["key"]), lang, custom),
                }
            )
        yield {"title": title, "lead": lead, "rows": rows}
        if scope != "step":
            for sub in step.get("subs") or []:
                yield {
                    "title": f"{sub['no']}  {_lang(sub['title'], lang)}",
                    "lead": "",
                    "rows": [
                        {
                            "label": _lang(f["label"], lang),
                            "hint": _lang(f["hint"], lang),
                            "value": _display(f, data.get(f["key"]), lang, custom),
                        }
                        for f in sub.get("fields") or []
                    ],
                }


def to_json_bytes(project: dict) -> bytes:
    return json.dumps(project, ensure_ascii=False, indent=2).encode("utf-8")


def to_docx_bytes(project: dict, scope: str, step_id: str | None, sub_id: str | None, lang: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    data = project.get("data") or project
    name = data.get("projectName") or PRODUCT_ZH
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    h = doc.add_paragraph(name)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = h.runs[0]
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x14, 0x08)

    sub = doc.add_paragraph(PRODUCT_ZH if lang == "zh" else PRODUCT_EN)
    sub.runs[0].font.size = Pt(11)
    sub.runs[0].font.color.rgb = RGBColor(0x6B, 0x5A, 0x3A)

    line = data.get("oneLiner") or ""
    if line:
        p = doc.add_paragraph(line)
        p.runs[0].italic = True

    for block in iter_sections(data, scope, step_id, sub_id, lang):
        ph = doc.add_heading(block["title"], level=1)
        if block["lead"]:
            lp = doc.add_paragraph(block["lead"])
            if lp.runs:
                lp.runs[0].font.size = Pt(10)
                lp.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        for row in block["rows"]:
            p = doc.add_paragraph()
            r = p.add_run(row["label"])
            r.bold = True
            r.font.size = Pt(12)
            if row["hint"]:
                hp = doc.add_paragraph(row["hint"])
                if hp.runs:
                    hp.runs[0].font.size = Pt(9)
                    hp.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            vp = doc.add_paragraph(row["value"])
            if vp.runs:
                vp.runs[0].font.size = Pt(11)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf_bytes(project: dict, scope: str, step_id: str | None, sub_id: str | None, lang: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "Helvetica"
    font_bold = "Helvetica-Bold"
    for path, alias in (
        ("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf", "NotoSC"),
        ("/usr/share/fonts/truetype/noto/NotoSansSC-Regular.ttf", "NotoSC"),
        ("C:/Windows/Fonts/msyh.ttc", "NotoSC"),
        ("C:/Windows/Fonts/simhei.ttf", "NotoSC"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "NotoSC"),
    ):
        from pathlib import Path

        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont(alias, path))
                font_name = alias
                font_bold = alias
                break
            except Exception:
                continue

    data = project.get("data") or project
    name = data.get("projectName") or PRODUCT_ZH
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=name,
        author=PRODUCT_EN,
    )
    styles = getSampleStyleSheet()
    title_s = ParagraphStyle(
        "T",
        parent=styles["Title"],
        fontName=font_bold,
        fontSize=18,
        textColor=HexColor("#1b1408"),
        spaceAfter=6,
        leading=24,
    )
    h_s = ParagraphStyle(
        "H",
        parent=styles["Heading1"],
        fontName=font_bold,
        fontSize=13,
        textColor=HexColor("#8a5a12"),
        spaceBefore=12,
        spaceAfter=4,
        leading=18,
    )
    hint_s = ParagraphStyle(
        "Hint",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8,
        textColor=HexColor("#888888"),
        leading=12,
        spaceAfter=2,
    )
    body_s = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=10,
        leading=15,
        spaceAfter=8,
    )
    label_s = ParagraphStyle(
        "Lab",
        parent=styles["Normal"],
        fontName=font_bold,
        fontSize=11,
        leading=15,
        spaceBefore=4,
        spaceAfter=1,
    )

    def esc(text: str) -> str:
        return (
            (text or "")
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )

    story = [Paragraph(esc(name), title_s)]
    story.append(Paragraph(esc(PRODUCT_ZH if lang == "zh" else PRODUCT_EN), hint_s))
    if data.get("oneLiner"):
        story.append(Paragraph(esc(str(data["oneLiner"])), body_s))
    story.append(HRFlowable(width="100%", color=HexColor("#e7c07a"), thickness=1, spaceAfter=8))

    for block in iter_sections(data, scope, step_id, sub_id, lang):
        story.append(Paragraph(esc(block["title"]), h_s))
        if block["lead"]:
            story.append(Paragraph(esc(block["lead"]), hint_s))
        for row in block["rows"]:
            story.append(Paragraph(esc(row["label"]), label_s))
            if row["hint"]:
                story.append(Paragraph(esc(row["hint"]), hint_s))
            story.append(Paragraph(esc(row["value"]), body_s))

    doc.build(story)
    return buf.getvalue()
